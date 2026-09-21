"""
Step 2: Read text out of resume files (PDF, DOCX, TXT).

Usage from the command line:
    python src/parser.py path/to/resume.pdf
"""
import io
import sys
from pathlib import Path


def parse_pdf(file_obj) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def parse_docx(file_obj) -> str:
    import docx

    document = docx.Document(file_obj)
    parts = [p.text for p in document.paragraphs]
    # Resumes often keep skills/experience inside tables
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def parse_txt(file_obj) -> str:
    data = file_obj.read()
    return data.decode("utf-8", errors="ignore") if isinstance(data, bytes) else data


def parse_resume(source, filename: str | None = None) -> str:
    """
    source   : a file path (str/Path) OR a file-like object (e.g. Streamlit upload)
    filename : needed only when `source` is a file-like object
    Returns the extracted plain text ('' if nothing could be read).
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = path.name
        file_obj = io.BytesIO(path.read_bytes())
    else:
        file_obj = io.BytesIO(source.read())

    ext = Path(filename or "").suffix.lower()
    file_obj.seek(0)

    if ext == ".pdf":
        return parse_pdf(file_obj)
    if ext == ".docx":
        return parse_docx(file_obj)
    if ext == ".txt":
        return parse_txt(file_obj)
    raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python src/parser.py <resume file>")
        sys.exit(1)
    text = parse_resume(sys.argv[1])
    print(f"Extracted {len(text.split())} words. First 500 characters:\n")
    print(text[:500])
